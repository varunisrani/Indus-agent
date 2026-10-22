# pyright: basic
"""Documents as plain text for ``read_file`` (P05-17).

``extract(data, name)`` picks a reader by extension and returns text that ``read_file`` pages with
its usual ``offset``/``limit``. Readers and their limits:

- Jupyter notebooks (stdlib JSON): each cell under a label; code-cell output text kept up to
  ``NOTEBOOK_OUTPUT_CHARS`` per cell, images, HTML and widgets replaced by a placeholder.
- PDF (``pypdf``): the text layer page by page; a page with fewer than ``PDF_EMPTY_PAGE_CHARS``
  characters counts as empty (likely a scan) and a note lists the empty page ranges (at most
  ``PDF_GAP_ENTRIES`` of them); the scan stops after ``PDF_SCAN_TIMEOUT_S``.
- Word ``.docx`` (``python-docx``): paragraphs and tables in document order.
- Excel ``.xlsx`` (``openpyxl``): every sheet, at most ``SHEET_MAX_ROWS`` rows and
  ``SHEET_MAX_COLS`` columns each.
- PowerPoint ``.pptx``, OpenDocument (``.odt``, ``.ods``, ``.odp``) and EPUB: zip files of XML,
  read with the standard library.
- RTF: our own control-word stripper.
- Zip-based formats (Word, Excel, PowerPoint, OpenDocument, EPUB) are checked before any reader
  opens them: at most ``ZIP_MAX_MEMBERS`` members whose declared sizes add up to at most
  ``ZIP_MAX_UNPACKED_BYTES``. ``zipfile`` never returns more than a member's declared size, so this
  bounds what every reader below decompresses (a small "zip bomb" cannot fill memory).
- SQLite (stdlib, read-only): each table's schema, row count and first ``SQLITE_PREVIEW_ROWS``
  rows, at most ``SQLITE_MAX_TABLES`` tables, cells cut at ``SQLITE_CELL_CHARS``.

PDF, Word and Excel need the optional ``docs`` extra; without it the reader raises
``MissingExtra`` naming the extra. Any failure is an ``ExtractionError`` with a plain reason.
"""

from __future__ import annotations

import contextlib
import html.parser
import io
import itertools
import json
import os
import posixpath
import re
import sqlite3
import tempfile
import zipfile
from collections.abc import Callable, Iterator
from typing import Any
from xml.etree import ElementTree

from indus.foundation.clock import Clock, SystemClock

__all__ = [
    "DOCUMENT_EXTENSIONS",
    "MAX_DOCUMENT_BYTES",
    "ExtractionError",
    "MissingExtra",
    "extract",
    "is_document",
]

MAX_DOCUMENT_BYTES = 50 * 1024 * 1024
NOTEBOOK_OUTPUT_CHARS = 20_000
PDF_EMPTY_PAGE_CHARS = 20
PDF_SCAN_TIMEOUT_S = 20.0
PDF_GAP_ENTRIES = 20
SHEET_MAX_ROWS = 5_000
SHEET_MAX_COLS = 256
SQLITE_PREVIEW_ROWS = 5
SQLITE_MAX_TABLES = 200
SQLITE_CELL_CHARS = 80
ZIP_MAX_MEMBERS = 10_000
ZIP_MAX_UNPACKED_BYTES = 256 * 1024 * 1024
INSTALL_HINT = "install the document readers with `uv sync --extra docs` (pip: indus-agent[docs])"


class ExtractionError(Exception):
    """The document could not be turned into text."""


class MissingExtra(ExtractionError):
    """The reader for this format lives in the ``docs`` extra, which is not installed."""


def _ext(name: str) -> str:
    base = posixpath.basename(name.replace("\\", "/")).lower()
    dot = base.rfind(".")
    return base[dot:] if dot > 0 else ""


def _finish(lines: list[str], empty: str) -> str:
    if not any(line.strip() for line in lines):
        raise ExtractionError(empty)
    return "\n".join(lines).rstrip("\n") + "\n"


def _missing(fmt: str, module: str) -> MissingExtra:
    return MissingExtra(f"reading {fmt} files needs the {module} package: {INSTALL_HINT}")


# -- notebooks -------------------------------------------------------------------------------------

_ANSI = re.compile(r"\x1b\[[0-9;?]*[A-Za-z]")


def _text_of(value: Any) -> str:
    if isinstance(value, list):
        return "".join(str(part) for part in value)
    return value if isinstance(value, str) else ""


def _output_text(output: Any) -> str:
    if not isinstance(output, dict):
        return ""
    kind = output.get("output_type")
    if kind == "stream":
        return _ANSI.sub("", _text_of(output.get("text")))
    if kind in ("error", "pyerr"):
        trace = output.get("traceback")
        lines = (
            [_ANSI.sub("", t) for t in trace if isinstance(t, str)]
            if isinstance(trace, list)
            else []
        )
        head = f"{output.get('ename', 'Error')}: {output.get('evalue', '')}".rstrip(": ")
        return "\n".join([head, *lines])
    if kind not in ("execute_result", "display_data", "pyout"):
        return ""
    data = output.get("data")
    if not isinstance(data, dict):
        data = {k: v for k, v in output.items() if k in ("text", "png", "jpeg", "html", "svg")}
        if "text" in data:
            data["text/plain"] = data.pop("text")
    if "application/vnd.jupyter.widget-view+json" in data:
        return "[interactive widget, not shown]"
    for mime in ("text/plain", "text/markdown"):
        body = _ANSI.sub("", _text_of(data.get(mime)))
        if body.strip():
            return body
    images = [
        m
        for m in data
        if isinstance(m, str) and (m.startswith("image/") or m in ("png", "jpeg", "svg"))
    ]
    if images:
        size = len(_text_of(data[images[0]])) * 3 // 4
        return f"[{images[0]} output, about {size // 1024} KB, not shown]"
    if "text/html" in data or "html" in data:
        return "[HTML output, not shown]"
    return f"[{', '.join(map(str, data)) or 'empty'} output, not shown]"


def _cell_outputs(cell: dict[str, Any]) -> str:
    outputs = cell.get("outputs")
    if not isinstance(outputs, list):
        return ""
    text = "\n".join(t.rstrip("\n") for t in map(_output_text, outputs) if t.strip())
    if len(text) <= NOTEBOOK_OUTPUT_CHARS:
        return text
    cut = len(text) - NOTEBOOK_OUTPUT_CHARS
    return text[:NOTEBOOK_OUTPUT_CHARS] + f"\n[... {cut:,} more output characters not shown]"


def _notebook(data: bytes) -> str:
    try:
        book = json.loads(data.decode("utf-8-sig", errors="replace"))
    except ValueError as exc:
        raise ExtractionError(f"not a valid notebook (JSON error: {exc})") from exc
    if not isinstance(book, dict):
        raise ExtractionError("not a valid notebook (the top level is not an object)")
    cells = book.get("cells")
    if not isinstance(cells, list):
        sheets = book.get("worksheets")
        cells = [
            c
            for sheet in (sheets if isinstance(sheets, list) else [])
            if isinstance(sheet, dict)
            for c in (sheet.get("cells") or [])
        ]
    lines: list[str] = []
    numbers: dict[str, int] = {}
    for cell in cells:
        if not isinstance(cell, dict):
            continue
        kind = cell.get("cell_type")
        if kind not in ("code", "markdown", "raw"):
            continue
        numbers[kind] = numbers.get(kind, 0) + 1
        source = _text_of(cell.get("source") or cell.get("input")).rstrip("\n")
        lines += [f"=== {kind} cell {numbers[kind]} ===", source, ""]
        if kind == "code":
            out = _cell_outputs(cell)
            if out:
                lines += [f"--- output of code cell {numbers[kind]} ---", out, ""]
    return _finish(lines, "the notebook has no cells")


# -- SQLite ----------------------------------------------------------------------------------------

_SQLITE_MAGIC = b"SQLite format 3\x00"


@contextlib.contextmanager
def _temp_copy(data: bytes, suffix: str) -> Iterator[str]:
    fd, path = tempfile.mkstemp(prefix="indus-doc-", suffix=suffix)
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(data)
        yield path
    finally:
        with contextlib.suppress(OSError):
            os.unlink(path)


def _cell(value: Any) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bytes | bytearray):
        return f"<{len(value)} bytes>"
    text = str(value).replace("\n", " ").replace("|", "\\|")
    return text if len(text) <= SQLITE_CELL_CHARS else text[: SQLITE_CELL_CHARS - 1] + "…"


def _quote_ident(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def _sqlite(data: bytes) -> str:
    if not data.startswith(_SQLITE_MAGIC):
        raise ExtractionError("not a SQLite database (the file header does not match)")
    with _temp_copy(data, ".db") as path:
        uri = "file:" + path + "?mode=ro&immutable=1"
        try:
            conn = sqlite3.connect(uri, uri=True)
        except sqlite3.Error as exc:
            raise ExtractionError(f"SQLite could not open the file: {exc}") from exc
        try:
            return _sqlite_summary(conn)
        except sqlite3.Error as exc:
            raise ExtractionError(f"SQLite could not read the file: {exc}") from exc
        finally:
            conn.close()


def _sqlite_summary(conn: sqlite3.Connection) -> str:
    rows = conn.execute(
        "SELECT type, name, sql FROM sqlite_master WHERE sql IS NOT NULL ORDER BY type, name"
    ).fetchall()
    tables = [(n, s) for t, n, s in rows if t == "table"]
    others = [(t, n) for t, n, _ in rows if t != "table"]
    lines = [
        "SQLite database",
        f"{len(tables)} table(s); {len(others)} index(es), view(s) or trigger(s)",
    ]
    for name, sql in tables[:SQLITE_MAX_TABLES]:
        ident = _quote_ident(name)
        count = conn.execute(f"SELECT COUNT(*) FROM {ident}").fetchone()[0]
        lines += ["", f"## table {name} ({count:,} rows)", sql.strip()]
        cursor = conn.execute(f"SELECT * FROM {ident} LIMIT {SQLITE_PREVIEW_ROWS}")
        columns = [d[0] for d in cursor.description or ()]
        preview = cursor.fetchall()
        if preview:
            lines.append("| " + " | ".join(columns) + " |")
            lines.extend("| " + " | ".join(_cell(v) for v in row) + " |" for row in preview)
    if len(tables) > SQLITE_MAX_TABLES:
        lines.append(f"\n[{len(tables) - SQLITE_MAX_TABLES} more tables not shown]")
    if others:
        lines += ["", "## other objects", *(f"- {kind} {name}" for kind, name in others)]
    lines += ["", "Query it through the terminal, e.g. sqlite3 <file> 'SELECT ...'."]
    return "\n".join(lines) + "\n"


# -- PDF -------------------------------------------------------------------------------------------


def _ranges(pages: list[int]) -> list[tuple[int, int]]:
    out: list[tuple[int, int]] = []
    for _, group in itertools.groupby(enumerate(pages), lambda pair: pair[1] - pair[0]):
        run = [p for _, p in group]
        out.append((run[0], run[-1]))
    return out


def _gap_note(empty: list[int], total: int) -> str:
    spans = _ranges(empty)
    shown = [f"page {a}" if a == b else f"pages {a}-{b}" for a, b in spans[:PDF_GAP_ENTRIES]]
    if len(spans) > PDF_GAP_ENTRIES:
        rest = spans[PDF_GAP_ENTRIES:]
        shown.append(f"... {len(rest)} more gaps ({sum(b - a + 1 for a, b in rest)} pages)")
    return (
        f"[{len(empty)} of {total} pages have no text layer (probably scanned images): "
        + ", ".join(shown)
        + ". Their content is missing below; render those pages to images and look at them "
        "if you need them.]"
    )


def _pdf(data: bytes, clock: Clock) -> str:
    try:
        from pypdf import PdfReader  # pyright: ignore[reportMissingImports] — optional extra
        from pypdf.errors import PdfReadError  # pyright: ignore[reportMissingImports]
    except ImportError:
        raise _missing("PDF", "pypdf") from None
    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted and not reader.decrypt(""):
            raise ExtractionError("the PDF is encrypted with a password")
        pages = list(reader.pages)
    except (PdfReadError, ValueError, OSError, KeyError) as exc:
        raise ExtractionError(f"not a readable PDF: {exc}") from exc
    if not pages:
        raise ExtractionError("the PDF has no pages")
    started = clock.monotonic()
    texts: list[str] = []
    stopped_at: int | None = None
    for number, page in enumerate(pages, start=1):
        if clock.monotonic() - started > PDF_SCAN_TIMEOUT_S:
            stopped_at = number
            break
        try:
            texts.append(page.extract_text() or "")
        except Exception as exc:  # pypdf raises many kinds on damaged pages
            texts.append(f"[page {number} could not be read: {type(exc).__name__}]")
    empty = [i for i, t in enumerate(texts, start=1) if len(t.strip()) < PDF_EMPTY_PAGE_CHARS]
    lines: list[str] = []
    if empty:
        lines += [_gap_note(empty, len(pages)), ""]
    for number, text in enumerate(texts, start=1):
        lines += [f"=== page {number} ===", text.rstrip("\n"), ""]
    if stopped_at is not None:
        lines.append(
            f"[stopped after {PDF_SCAN_TIMEOUT_S:.0f} s at page {stopped_at} of {len(pages)}; "
            "split the PDF or extract the rest through the terminal]"
        )
    return "\n".join(lines).rstrip("\n") + "\n"


# -- Word and Excel --------------------------------------------------------------------------------


def _docx(data: bytes) -> str:
    try:
        import docx  # pyright: ignore[reportMissingImports] — optional extra
        from docx.table import Table  # pyright: ignore[reportMissingImports]
    except ImportError:
        raise _missing("Word (.docx)", "python-docx") from None
    try:
        document = docx.Document(io.BytesIO(data))
        lines: list[str] = []
        for block in document.iter_inner_content():
            if isinstance(block, Table):
                lines.append("")
                for row in block.rows:
                    lines.append(" | ".join(cell.text.replace("\n", " ") for cell in row.cells))
                lines.append("")
            else:
                lines.append(block.text)
    except (KeyError, ValueError, zipfile.BadZipFile) as exc:
        raise ExtractionError(f"not a readable Word document: {exc}") from exc
    return _finish(lines, "the Word document has no text")


def _xlsx(data: bytes) -> str:
    try:
        import openpyxl  # pyright: ignore[reportMissingModuleSource]
    except ImportError:
        raise _missing("Excel (.xlsx)", "openpyxl") from None
    try:
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (KeyError, ValueError, OSError, zipfile.BadZipFile) as exc:
        raise ExtractionError(f"not a readable Excel workbook: {exc}") from exc
    lines: list[str] = []
    try:
        for sheet in book.worksheets:
            lines.append(f"## sheet {sheet.title}")
            wide = False
            shown = 0
            more = 0
            for row in sheet.iter_rows(values_only=True):
                if shown >= SHEET_MAX_ROWS:
                    more += 1
                    continue
                cells = list(row)
                if len(cells) > SHEET_MAX_COLS:
                    wide = True
                    cells = cells[:SHEET_MAX_COLS]
                while cells and cells[-1] is None:
                    cells.pop()
                lines.append(" | ".join("" if v is None else str(v) for v in cells))
                shown += 1
            if more:
                lines.append(f"[{more:,} more rows not shown (limit {SHEET_MAX_ROWS:,})]")
            if wide:
                lines.append(f"[columns after {SHEET_MAX_COLS} not shown]")
            lines.append("")
    finally:
        book.close()
    return _finish(lines, "the workbook has no sheets")


# -- zip + XML formats -----------------------------------------------------------------------------

_ZIP_FORMATS = frozenset({".docx", ".xlsx", ".pptx", ".odt", ".ods", ".odp", ".epub"})


def _check_zip_budget(data: bytes) -> None:
    """Refuse a zip that would unpack past the budget (see the module notes). A file that is not
    a zip passes: its reader reports that in its own words."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
    except (zipfile.BadZipFile, ValueError, OSError):
        return
    if len(members) > ZIP_MAX_MEMBERS:
        raise ExtractionError(
            f"the document holds {len(members):,} parts; more than {ZIP_MAX_MEMBERS:,} are not read"
        )
    unpacked = sum(max(0, m.file_size) for m in members)
    if unpacked > ZIP_MAX_UNPACKED_BYTES:
        raise ExtractionError(
            f"the document would unpack to {unpacked / 1048576:,.0f} MB; more than "
            f"{ZIP_MAX_UNPACKED_BYTES // 1048576} MB is not extracted"
        )


def _open_zip(data: bytes, what: str) -> zipfile.ZipFile:
    try:
        return zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ExtractionError(f"not a readable {what} (bad zip: {exc})") from exc


def _xml(archive: zipfile.ZipFile, name: str) -> ElementTree.Element:
    try:
        return ElementTree.fromstring(archive.read(name))
    except KeyError as exc:
        raise ExtractionError(f"missing part {name}") from exc
    except ElementTree.ParseError as exc:
        raise ExtractionError(f"broken XML in {name}: {exc}") from exc


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _paragraph_texts(root: ElementTree.Element, para: str, text: str) -> list[str]:
    out: list[str] = []
    for node in root.iter():
        if _local(node.tag) == para:
            out.append("".join(t.text or "" for t in node.iter() if _local(t.tag) == text))
    return out


def _slide_number(name: str) -> int:
    match = re.search(r"(\d+)\.xml$", name)
    return int(match.group(1)) if match else 0


def _pptx(data: bytes) -> str:
    with _open_zip(data, "PowerPoint file") as archive:
        slides = sorted(
            (n for n in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
            key=_slide_number,
        )
        lines: list[str] = []
        for name in slides:
            lines.append(f"=== slide {_slide_number(name)} ===")
            lines += [t for t in _paragraph_texts(_xml(archive, name), "p", "t") if t.strip()]
            lines.append("")
    return _finish(lines, "the presentation has no text")


def _attr(node: ElementTree.Element, name: str, default: str = "") -> str:
    """An attribute by its local name, whatever its namespace."""
    for key, value in node.attrib.items():
        if _local(key) == name:
            return value
    return default


def _odf_text(node: ElementTree.Element) -> str:
    parts = [node.text or ""]
    for child in node:
        name = _local(child.tag)
        if name == "s":
            count = _attr(child, "c", "1")
            parts.append(" " * (int(count) if count.isdigit() else 1))
        elif name == "tab":
            parts.append("\t")
        elif name == "line-break":
            parts.append("\n")
        else:
            parts.append(_odf_text(child))
        parts.append(child.tail or "")
    return "".join(parts)


def _odf(data: bytes) -> str:
    with _open_zip(data, "OpenDocument file") as archive:
        root = _xml(archive, "content.xml")
    lines: list[str] = []
    for table in (n for n in root.iter() if _local(n.tag) == "table"):
        lines.append(f"## table {_attr(table, 'name')}".rstrip())
        for count, row in enumerate(n for n in table.iter() if _local(n.tag) == "table-row"):
            if count >= SHEET_MAX_ROWS:
                lines.append(f"[rows after {SHEET_MAX_ROWS:,} not shown]")
                break
            cells = [
                " ".join(_odf_text(p) for p in cell.iter() if _local(p.tag) == "p")
                for cell in row
                if _local(cell.tag) == "table-cell"
            ][:SHEET_MAX_COLS]
            while cells and not cells[-1]:
                cells.pop()
            if cells:
                lines.append(" | ".join(cells))
        lines.append("")
    in_table = {id(p) for t in root.iter() if _local(t.tag) == "table" for p in t.iter()}
    for node in root.iter():
        if _local(node.tag) in ("h", "p") and id(node) not in in_table:
            lines.append(_odf_text(node))
    return _finish(lines, "the document has no text")


class _HtmlText(html.parser.HTMLParser):
    _BLOCKS = frozenset(
        {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section"}
    )
    _SKIP = frozenset({"script", "style", "head"})

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skipping = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._SKIP:
            self._skipping += 1
        elif tag in self._BLOCKS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP and self._skipping:
            self._skipping -= 1
        elif tag in self._BLOCKS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skipping:
            self.parts.append(data)

    def text(self) -> str:
        joined = "".join(self.parts)
        lines = [" ".join(line.split()) for line in joined.splitlines()]
        return "\n".join(line for line in lines if line)


def _epub(data: bytes) -> str:
    with _open_zip(data, "EPUB book") as archive:
        container = _xml(archive, "META-INF/container.xml")
        rootfile = next(
            (n.get("full-path") for n in container.iter() if _local(n.tag) == "rootfile"), None
        )
        if not rootfile:
            raise ExtractionError("the EPUB has no package file")
        package = _xml(archive, rootfile)
        base = posixpath.dirname(rootfile)
        manifest = {
            n.get("id"): n.get("href")
            for n in package.iter()
            if _local(n.tag) == "item" and n.get("id") and n.get("href")
        }
        order = [n.get("idref") for n in package.iter() if _local(n.tag) == "itemref"]
        lines: list[str] = []
        for chapter, ref in enumerate(order, start=1):
            href = manifest.get(ref)
            if not href:
                continue
            try:
                raw = archive.read(posixpath.normpath(posixpath.join(base, href)))
            except KeyError:
                continue
            parser = _HtmlText()
            parser.feed(raw.decode("utf-8", errors="replace"))
            lines += [f"=== section {chapter} ===", parser.text(), ""]
    return _finish(lines, "the book has no text")


# -- RTF -------------------------------------------------------------------------------------------

_RTF_TOKEN = re.compile(
    r"\\([a-zA-Z]+)(-?\d+)? ?|\\'([0-9a-fA-F]{2})|\\(.)|([{}])|([^\\{}]+)", re.S
)
_RTF_SKIP = frozenset(
    {
        "fonttbl",
        "colortbl",
        "stylesheet",
        "info",
        "pict",
        "header",
        "footer",
        "headerl",
        "headerr",
        "footerl",
        "footerr",
        "themedata",
        "colorschememapping",
        "latentstyles",
        "datastore",
        "xmlnstbl",
        "listtable",
        "listoverridetable",
        "rsidtbl",
        "generator",
    }
)
_RTF_BREAKS = {
    "par": "\n",
    "line": "\n",
    "tab": "\t",
    "emdash": "—",
    "endash": "–",
    "bullet": "•",
    "lquote": "‘",
    "rquote": "’",
    "ldblquote": "“",
    "rdblquote": "”",
}


def _rtf(data: bytes) -> str:
    text = data.decode("latin-1")
    if not text.lstrip().startswith("{\\rtf"):
        raise ExtractionError("not an RTF file (it does not start with {\\rtf)")
    out: list[str] = []
    stack: list[bool] = []
    skipping = False
    group_start = False
    uc_skip = 1
    pending_skip = 0
    for word, arg, hexbyte, symbol, brace, plain in _RTF_TOKEN.findall(text):
        if brace == "{":
            stack.append(skipping)
            group_start = True
            continue
        if brace == "}":
            skipping = stack.pop() if stack else False
            group_start = False
            continue
        starting, group_start = group_start, False
        if symbol == "*" and starting:
            skipping = True
            continue
        if word:
            if starting and word in _RTF_SKIP:
                skipping = True
            elif word == "uc" and arg:
                uc_skip = int(arg)
            elif word == "u" and arg and not skipping:
                code = int(arg)
                out.append(chr(code + 65536 if code < 0 else code))
                pending_skip = uc_skip
            elif not skipping and word in _RTF_BREAKS:
                out.append(_RTF_BREAKS[word])
            continue
        if pending_skip:
            if hexbyte:
                pending_skip -= 1
                continue
            if plain:
                plain = plain[pending_skip:]
                pending_skip = 0
        if skipping:
            continue
        if hexbyte:
            out.append(bytes([int(hexbyte, 16)]).decode("cp1252", errors="replace"))
        elif symbol:
            out.append({"~": "\u00a0", "-": "", "_": "-"}.get(symbol, symbol))
        elif plain:
            out.append(plain.replace("\r", "").replace("\n", ""))
    lines = "".join(out).splitlines()
    return _finish([line.rstrip() for line in lines], "the RTF file has no text")


# -- the table -------------------------------------------------------------------------------------

_READERS: dict[str, Callable[[bytes, Clock], str]] = {
    ".ipynb": lambda data, _clock: _notebook(data),
    ".pdf": _pdf,
    ".docx": lambda data, _clock: _docx(data),
    ".xlsx": lambda data, _clock: _xlsx(data),
    ".pptx": lambda data, _clock: _pptx(data),
    ".odt": lambda data, _clock: _odf(data),
    ".ods": lambda data, _clock: _odf(data),
    ".odp": lambda data, _clock: _odf(data),
    ".epub": lambda data, _clock: _epub(data),
    ".rtf": lambda data, _clock: _rtf(data),
    ".db": lambda data, _clock: _sqlite(data),
    ".sqlite": lambda data, _clock: _sqlite(data),
    ".sqlite3": lambda data, _clock: _sqlite(data),
}
DOCUMENT_EXTENSIONS = frozenset(_READERS)


def is_document(name: str) -> bool:
    return _ext(name) in _READERS


def extract(data: bytes, name: str, *, clock: Clock | None = None) -> str:
    """``data`` (the bytes of the file called ``name``) as text; see the module notes."""
    reader = _READERS.get(_ext(name))
    if reader is None:
        raise ExtractionError(f"no reader for {_ext(name) or 'files without an extension'}")
    if len(data) > MAX_DOCUMENT_BYTES:
        raise ExtractionError(
            f"the document is {len(data) / 1048576:.1f} MB; documents over "
            f"{MAX_DOCUMENT_BYTES // 1048576} MB are not extracted"
        )
    if _ext(name) in _ZIP_FORMATS:
        _check_zip_budget(data)
    try:
        return reader(data, clock if clock is not None else SystemClock())
    except ExtractionError:
        raise
    except RecursionError as exc:
        raise ExtractionError("the document is nested too deeply to read") from exc
    except Exception as exc:  # third-party readers raise many kinds on damaged files
        raise ExtractionError(
            f"the document could not be read: {type(exc).__name__}: {exc}"
        ) from exc
